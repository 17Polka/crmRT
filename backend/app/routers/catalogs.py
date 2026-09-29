"""
routers/catalogs.py — CRUD-роутер для раздела «Каталоги данных».

Эндпоинты:
  Directions  GET/POST /api/catalogs/directions
              GET/PUT/DELETE /api/catalogs/directions/{id}
  Products    GET/POST /api/catalogs/products
              GET/PUT/DELETE /api/catalogs/products/{id}
  Universities GET/POST /api/catalogs/universities
               GET/PUT/DELETE /api/catalogs/universities/{id}
  Import      POST /api/catalogs/import/universities  (xlsx)

Права:
  • Чтение справочников — любой авторизованный (manager+)
  • Запись/изменение — head+
  • Удаление / импорт — admin
"""
import io
import json
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import AuditAction, Direction, Product, University, User
from app.schemas import (
    DirectionCreate,
    DirectionOut,
    DirectionUpdate,
    ImportResult,
    MessageResponse,
    PagedResponse,
    ProductCreate,
    ProductOut,
    ProductUpdate,
    UniversityCreate,
    UniversityOut,
    UniversityUpdate,
)
from app.security import require_admin, require_head, require_manager, write_audit

router = APIRouter(prefix="/api/catalogs", tags=["Каталоги"])


# ===========================================================================
# DIRECTIONS
# ===========================================================================

@router.get("/directions", response_model=PagedResponse, summary="Список направлений")
async def list_directions(
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    q: Optional[str] = None,
    is_active: Optional[bool] = None,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    stmt = select(Direction)
    if q:
        stmt = stmt.where(Direction.name.ilike(f"%{q}%"))
    if is_active is not None:
        stmt = stmt.where(Direction.is_active == is_active)

    total = await db.scalar(select(func.count()).select_from(stmt.subquery()))
    items = (
        await db.scalars(
            stmt.order_by(Direction.name)
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
    ).all()

    return PagedResponse(
        total=total or 0,
        page=page,
        page_size=page_size,
        items=[DirectionOut.model_validate(d) for d in items],
    )


@router.post("/directions", response_model=DirectionOut, status_code=201,
             summary="Создать направление")
async def create_direction(
    body: DirectionCreate,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_head),
):
    # Уникальность
    exists = await db.scalar(
        select(Direction).where(Direction.name == body.name)
    )
    if exists:
        raise HTTPException(400, detail=f"Направление '{body.name}' уже существует.")

    obj = Direction(**body.model_dump())
    db.add(obj)
    await db.flush()  # получаем id до audit

    await write_audit(
        db, action=AuditAction.catalog_created,
        user=current_user, resource_type="direction", resource_id=obj.id,
        detail={"name": obj.name}, request=request,
    )
    await db.commit()
    await db.refresh(obj)
    return DirectionOut.model_validate(obj)


@router.get("/directions/{direction_id}", response_model=DirectionOut,
            summary="Получить направление")
async def get_direction(
    direction_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    obj = await db.get(Direction, direction_id)
    if not obj:
        raise HTTPException(404, detail="Направление не найдено.")
    return DirectionOut.model_validate(obj)


@router.put("/directions/{direction_id}", response_model=DirectionOut,
            summary="Обновить направление")
async def update_direction(
    direction_id: int,
    body: DirectionUpdate,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_head),
):
    obj = await db.get(Direction, direction_id)
    if not obj:
        raise HTTPException(404, detail="Направление не найдено.")

    changes = body.model_dump(exclude_unset=True)
    for k, v in changes.items():
        setattr(obj, k, v)

    await write_audit(
        db, action=AuditAction.catalog_updated,
        user=current_user, resource_type="direction", resource_id=direction_id,
        detail=changes, request=request,
    )
    await db.commit()
    await db.refresh(obj)
    return DirectionOut.model_validate(obj)


@router.delete("/directions/{direction_id}", response_model=MessageResponse,
               summary="Удалить направление")
async def delete_direction(
    direction_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    obj = await db.get(Direction, direction_id)
    if not obj:
        raise HTTPException(404, detail="Направление не найдено.")
    # Мягкое удаление — сохраняем данные
    obj.is_active = False
    await write_audit(
        db, action=AuditAction.catalog_deleted,
        user=current_user, resource_type="direction", resource_id=direction_id,
        request=request,
    )
    await db.commit()
    return MessageResponse(ok=True, message=f"Направление #{direction_id} деактивировано.")


# ===========================================================================
# PRODUCTS
# ===========================================================================

@router.get("/products", response_model=PagedResponse, summary="Список продуктов")
async def list_products(
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    q: Optional[str] = None,
    direction_id: Optional[int] = None,
    is_active: Optional[bool] = None,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    stmt = select(Product)
    if q:
        stmt = stmt.where(Product.name.ilike(f"%{q}%"))
    if direction_id:
        stmt = stmt.where(Product.direction_id == direction_id)
    if is_active is not None:
        stmt = stmt.where(Product.is_active == is_active)

    total = await db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = (
        await db.scalars(
            stmt.order_by(Product.name)
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
    ).all()

    # Подтягиваем имя направления
    dir_cache: dict[int, str] = {}
    items_out = []
    for p in rows:
        if p.direction_id not in dir_cache:
            d = await db.get(Direction, p.direction_id)
            dir_cache[p.direction_id] = d.name if d else ""
        out = ProductOut.model_validate(p)
        out.direction_name = dir_cache[p.direction_id]
        items_out.append(out)

    return PagedResponse(total=total or 0, page=page, page_size=page_size, items=items_out)


@router.post("/products", response_model=ProductOut, status_code=201,
             summary="Создать продукт")
async def create_product(
    body: ProductCreate,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_head),
):
    # Проверяем направление
    direction = await db.get(Direction, body.direction_id)
    if not direction or not direction.is_active:
        raise HTTPException(400, detail="Направление не найдено или неактивно.")

    exists = await db.scalar(
        select(Product).where(Product.name == body.name)
    )
    if exists:
        raise HTTPException(400, detail=f"Продукт '{body.name}' уже существует.")

    obj = Product(**body.model_dump())
    db.add(obj)
    await db.flush()

    await write_audit(
        db, action=AuditAction.catalog_created,
        user=current_user, resource_type="product", resource_id=obj.id,
        detail={"name": obj.name}, request=request,
    )
    await db.commit()
    await db.refresh(obj)

    out = ProductOut.model_validate(obj)
    out.direction_name = direction.name
    return out


@router.get("/products/{product_id}", response_model=ProductOut, summary="Получить продукт")
async def get_product(
    product_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    obj = await db.get(Product, product_id)
    if not obj:
        raise HTTPException(404, detail="Продукт не найден.")
    direction = await db.get(Direction, obj.direction_id)
    out = ProductOut.model_validate(obj)
    out.direction_name = direction.name if direction else None
    return out


@router.put("/products/{product_id}", response_model=ProductOut, summary="Обновить продукт")
async def update_product(
    product_id: int,
    body: ProductUpdate,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_head),
):
    obj = await db.get(Product, product_id)
    if not obj:
        raise HTTPException(404, detail="Продукт не найден.")

    changes = body.model_dump(exclude_unset=True)
    if "direction_id" in changes:
        direction = await db.get(Direction, changes["direction_id"])
        if not direction or not direction.is_active:
            raise HTTPException(400, detail="Указанное направление не найдено или неактивно.")

    for k, v in changes.items():
        setattr(obj, k, v)

    await write_audit(
        db, action=AuditAction.catalog_updated,
        user=current_user, resource_type="product", resource_id=product_id,
        detail=changes, request=request,
    )
    await db.commit()
    await db.refresh(obj)

    direction = await db.get(Direction, obj.direction_id)
    out = ProductOut.model_validate(obj)
    out.direction_name = direction.name if direction else None
    return out


@router.delete("/products/{product_id}", response_model=MessageResponse,
               summary="Удалить продукт")
async def delete_product(
    product_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    obj = await db.get(Product, product_id)
    if not obj:
        raise HTTPException(404, detail="Продукт не найден.")
    obj.is_active = False
    await write_audit(
        db, action=AuditAction.catalog_deleted,
        user=current_user, resource_type="product", resource_id=product_id,
        request=request,
    )
    await db.commit()
    return MessageResponse(ok=True, message=f"Продукт #{product_id} деактивирован.")


# ===========================================================================
# UNIVERSITIES (справочник)
# ===========================================================================

@router.get("/universities", response_model=PagedResponse, summary="Список вузов (справочник)")
async def list_universities(
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    q: Optional[str] = None,
    city: Optional[str] = None,
    direction_id: Optional[int] = None,
    product_id: Optional[int] = None,
    is_active: Optional[bool] = True,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_manager),
):
    stmt = select(University)
    if q:
        stmt = stmt.where(
            University.name.ilike(f"%{q}%") | University.short_name.ilike(f"%{q}%")
        )
    if city:
        stmt = stmt.where(University.city.ilike(f"%{city}%"))
    if direction_id:
        stmt = stmt.where(University.direction_id == direction_id)
    if product_id:
        stmt = stmt.where(University.product_id == product_id)
    if is_active is not None:
        stmt = stmt.where(University.is_active == is_active)

    # Менеджер видит только свои вузы
    from app.models import UserRole
    if current_user.role == UserRole.manager:
        stmt = stmt.where(University.manager_id == current_user.id)

    total = await db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = (
        await db.scalars(
            stmt.order_by(University.name)
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
    ).all()

    items_out = []
    for u in rows:
        out = UniversityOut.model_validate(u)
        if u.direction_id:
            d = await db.get(Direction, u.direction_id)
            out.direction_name = d.name if d else None
        if u.product_id:
            p = await db.get(Product, u.product_id)
            out.product_name = p.name if p else None
        if u.manager_id:
            m = await db.get(User, u.manager_id)
            out.manager_name = m.full_name if m else None
        items_out.append(out)

    return PagedResponse(total=total or 0, page=page, page_size=page_size, items=items_out)


@router.post("/universities", response_model=UniversityOut, status_code=201,
             summary="Добавить вуз в справочник")
async def create_university(
    body: UniversityCreate,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_head),
):
    exists = await db.scalar(select(University).where(University.name == body.name))
    if exists:
        raise HTTPException(400, detail=f"Вуз '{body.name}' уже существует.")

    obj = University(**body.model_dump())
    db.add(obj)
    await db.flush()

    await write_audit(
        db, action=AuditAction.catalog_created,
        user=current_user, resource_type="university", resource_id=obj.id,
        detail={"name": obj.name}, request=request,
    )
    await db.commit()
    await db.refresh(obj)
    return UniversityOut.model_validate(obj)


@router.get("/universities/{uni_id}", response_model=UniversityOut,
            summary="Получить вуз")
async def get_university(
    uni_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_manager),
):
    obj = await db.get(University, uni_id)
    if not obj:
        raise HTTPException(404, detail="Вуз не найден.")
    # Менеджер видит только свои
    from app.models import UserRole
    if current_user.role == UserRole.manager and obj.manager_id != current_user.id:
        raise HTTPException(403, detail="Нет доступа к этому вузу.")
    return UniversityOut.model_validate(obj)


@router.put("/universities/{uni_id}", response_model=UniversityOut,
            summary="Обновить вуз")
async def update_university(
    uni_id: int,
    body: UniversityUpdate,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_head),
):
    obj = await db.get(University, uni_id)
    if not obj:
        raise HTTPException(404, detail="Вуз не найден.")

    changes = body.model_dump(exclude_unset=True)
    for k, v in changes.items():
        setattr(obj, k, v)

    if "manager_id" in changes and changes["manager_id"]:
        mgr = await db.get(User, changes["manager_id"])
        if mgr:
            obj.manager_fio = mgr.full_name

    await write_audit(
        db, action=AuditAction.catalog_updated,
        user=current_user, resource_type="university", resource_id=uni_id,
        detail=changes, request=request,
    )
    await db.commit()
    await db.refresh(obj)
    return UniversityOut.model_validate(obj)


@router.delete("/universities/{uni_id}", response_model=MessageResponse,
               summary="Удалить вуз из справочника")
async def delete_university(
    uni_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    obj = await db.get(University, uni_id)
    if not obj:
        raise HTTPException(404, detail="Вуз не найден.")
    obj.is_active = False
    await write_audit(
        db, action=AuditAction.catalog_deleted,
        user=current_user, resource_type="university", resource_id=uni_id,
        request=request,
    )
    await db.commit()
    return MessageResponse(ok=True, message=f"Вуз #{uni_id} деактивирован.")


# ===========================================================================
# IMPORT — xlsx/xls → university catalog (Поддержка всех полей ТЗ и синонимов)
# ===========================================================================

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

HEADER_ALIASES: dict[str, list[str]] = {
    "name": [
        "name", "название", "наименование", "наименование вуза", "вуз",
        "полное наименование", "учебное заведение", "университет", "организация",
        "название вуза", "наименование организации"
    ],
    "short_name": [
        "short_name", "краткое", "краткое наименование", "краткое название",
        "аббревиатура", "сокр", "короткое имя"
    ],
    "city": [
        "city", "город", "регион", "населенный пункт", "субъект рф", "субъект", "местонахождение"
    ],
    "inn": [
        "inn", "инн", "идентификационный номер", "инн вуза"
    ],
    "licence_year": [
        "licence_year", "license_year", "срок лицензии", "год лицензии",
        "лицензия", "срок действия", "лицензия до", "год"
    ],
    "direction": [
        "direction", "направление", "ит-направление", "программа",
        "программа обучения", "направление подготовки"
    ],
    "product": [
        "product", "продукт", "ит-продукт", "по", "софт", "вендорский продукт", "программное обеспечение"
    ],
    "manager": [
        "manager", "менеджер", "кам", "ответственный", "курирующий менеджер", "ответственный кам"
    ],
    "contract": [
        "contract", "договор", "номер договора", "соглашение", "номер соглашения"
    ],
    "contacts": [
        "contacts", "контакты", "представители", "контактное лицо", "телефон", "email", "контактные данные"
    ],
}


def _read_spreadsheet_rows(content: bytes, filename: str) -> list[list]:
    lower_name = filename.lower()
    if lower_name.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        return list(ws.iter_rows(values_only=True))
    elif lower_name.endswith(".xls"):
        if HAS_XLRD:
            try:
                wb = xlrd.open_workbook(file_contents=content)
                ws = wb.sheet_by_index(0)
                rows = []
                for r in range(ws.nrows):
                    rows.append([ws.cell_value(r, c) for c in range(ws.ncols)])
                return rows
            except Exception:
                pass
        # Fallback на случай, если .xlsx сохранили с расширением .xls
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            return list(ws.iter_rows(values_only=True))
        except Exception:
            raise HTTPException(
                400,
                detail="Не удалось открыть файл .xls. Рекомендуется использовать современный формат .xlsx."
            )
    else:
        raise HTTPException(400, detail="Поддерживаются только форматы .xlsx и .xls")


@router.post(
    "/import/universities",
    response_model=ImportResult,
    summary="Импорт вузов из Excel (.xlsx / .xls)",
    description=(
        "Поддерживает как английские (`name`, `inn`), так и русскоязычные заголовки "
        "('Наименование ВУЗа', 'ИНН', 'Город', 'ИТ-Направление', 'ИТ-Продукт', 'Срок лицензии', 'Ответственный КАМ', 'Номер договора')."
    ),
)
async def import_universities_xlsx(
    request: Request,
    file: UploadFile = File(..., description="Excel-файл со справочником вузов (.xlsx, .xls)"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, detail="Ожидается файл .xlsx или .xls")

    content = await file.read()
    try:
        rows = _read_spreadsheet_rows(content, file.filename)
    except Exception as exc:
        if isinstance(exc, HTTPException):
            raise exc
        raise HTTPException(400, detail=f"Ошибка чтения таблицы: {exc}")

    if not rows or len(rows) < 2:
        raise HTTPException(400, detail="Файл пустой или содержит только строку заголовков.")

    # Нормализуем заголовки и строим маппинг колонок по синонимам
    raw_headers = [str(c).strip().lower() if c else "" for c in rows[0]]
    header_map: dict[str, int] = {}

    for col_key, aliases in HEADER_ALIASES.items():
        for col_idx, h_name in enumerate(raw_headers):
            if h_name in aliases or any(alias in h_name for alias in aliases):
                header_map[col_key] = col_idx
                break

    # Если 'name' не найден по алиасам, проверим первую колонку
    if "name" not in header_map:
        if raw_headers and raw_headers[0]:
            header_map["name"] = 0
        else:
            raise HTTPException(
                400,
                detail=f"Не найден столбец с наименованием ВУЗа (ожидались: 'Наименование', 'ВУЗ', 'name'). "
                       f"Обнаруженные заголовки: {raw_headers}",
            )

    # Предзагружаем справочники направлений, продуктов и менеджеров для связи
    existing_dirs = {d.name.lower(): d.id for d in (await db.scalars(select(Direction))).all()}
    existing_prods = {p.name.lower(): p.id for p in (await db.scalars(select(Product))).all()}
    managers = (await db.scalars(select(User).where(User.is_active == True))).all()

    created = updated = skipped = 0
    errors: list[str] = []

    for row_num, row in enumerate(rows[1:], start=2):
        if not row:
            skipped += 1
            continue

        def cell(col: str):
            idx = header_map.get(col)
            if idx is not None and idx < len(row):
                val = row[idx]
                return str(val).strip() if val is not None else None
            return None

        name = cell("name")
        if not name:
            skipped += 1
            continue

        try:
            inn = cell("inn")
            # Поиск по ИНН или по наименованию
            existing = None
            if inn:
                existing = await db.scalar(select(University).where(University.inn == inn))
            if not existing:
                existing = await db.scalar(select(University).where(University.name == name))

            # Связь с направлением
            dir_id = None
            dir_name = cell("direction")
            if dir_name:
                dir_id = existing_dirs.get(dir_name.lower())
                if not dir_id:
                    new_d = Direction(name=dir_name)
                    db.add(new_d)
                    await db.flush()
                    dir_id = new_d.id
                    existing_dirs[dir_name.lower()] = dir_id

            # Связь с продуктом
            prod_id = None
            prod_name = cell("product")
            if prod_name:
                prod_id = existing_prods.get(prod_name.lower())
                if not prod_id:
                    new_p = Product(name=prod_name, direction_id=dir_id)
                    db.add(new_p)
                    await db.flush()
                    prod_id = new_p.id
                    existing_prods[prod_name.lower()] = prod_id

            # Связь с менеджером
            mgr_id = None
            mgr_str = cell("manager")
            if mgr_str:
                m_lower = mgr_str.lower()
                for m in managers:
                    if m.username.lower() == m_lower or m_lower in m.full_name.lower():
                        mgr_id = m.id
                        break

            # Преобразование года лицензии
            lic_year = None
            lic_val = cell("licence_year")
            if lic_val:
                try:
                    lic_year = int(float(lic_val))
                except (ValueError, TypeError):
                    pass

            data = {
                "name": name,
                "short_name": cell("short_name"),
                "city": cell("city"),
                "inn": inn,
                "licence_year": lic_year,
                "contract_number": cell("contract"),
                "university_contacts": cell("contacts"),
            }
            if dir_id:
                data["direction_id"] = dir_id
            if prod_id:
                data["product_id"] = prod_id
            if mgr_id:
                data["manager_id"] = mgr_id

            if existing:
                for k, v in data.items():
                    if v is not None:
                        setattr(existing, k, v)
                updated += 1
            else:
                obj = University(**data)
                db.add(obj)
                created += 1

        except Exception as exc:
            errors.append(f"Строка {row_num}: {exc}")
            skipped += 1

    await write_audit(
        db, action=AuditAction.catalog_imported,
        user=current_user, resource_type="university",
        detail={"file": file.filename, "created": created, "updated": updated,
                "skipped": skipped, "errors": len(errors)},
        request=request,
    )
    await db.commit()

    return ImportResult(
        ok=True, created=created, updated=updated,
        skipped=skipped, errors=errors,
    )


@router.get("/managers", summary="Список ответственных менеджеров")
async def list_managers(
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    stmt = (
        select(User)
        .where(User.is_active == True, User.is_blocked == False)
        .order_by(User.full_name)
    )
    users = (await db.scalars(stmt)).all()
    return [
        {
            "id": u.id,
            "username": u.username,
            "full_name": u.full_name,
            "role": u.role.value,
            "email": u.email,
        }
        for u in users
    ]
