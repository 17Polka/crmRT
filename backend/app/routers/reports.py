"""
reports.py — Формирование отчётов в форматах XLSX, PDF и JSON.
Соответствие требованиям ТЗ (п. 2, 4, 8 раздела «Требования к сервису»):
- Выгрузка за выбранный период
- Фильтрация по вузам, направлениям, продуктам, ответственным
- Выбор колонок: наименование вуза, ИТ-направление, ИТ-продукт, статус работы, ответственный
- Форматы: xlsx, pdf, json (результирующий JSON п. 6.4)
"""
import io
import json
from datetime import datetime
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy import select, or_
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from app.database import get_db
from app.models import University, Direction, Product, User, WorkflowStage, UserRole, WorkflowHistory
from app.security import get_current_user, require_manager

router = APIRouter(prefix="/api/reports", tags=["Отчёты"])


# Регистрируем шрифт для кириллицы в PDF если есть Arial в Windows, иначе fallback
try:
    pdfmetrics.registerFont(TTFont("Arial", "C:\\Windows\\Fonts\\arial.ttf"))
    PDF_FONT = "Arial"
except Exception:
    PDF_FONT = "Helvetica"


COLUMN_MAP = {
    "name": ("ВУЗ", lambda u, d, p: u.name),
    "city": ("Город", lambda u, d, p: u.city or "—"),
    "direction": ("ИТ-Направление", lambda u, d, p: d.name if d else "—"),
    "product": ("ИТ-Продукт", lambda u, d, p: p.name if p else (u.software or "—")),
    "vendor": ("Вендор", lambda u, d, p: u.vendor or (p.vendor if p else "—")),
    "stage": ("Статус работы", lambda u, d, p: f"Этап {u.current_stage_order + 1}"),
    "manager": ("Ответственный", lambda u, d, p: u.manager_fio or "Не назначен"),
    "contract": ("Номер договора", lambda u, d, p: u.contract_number or "—"),
    "licence_year": ("Срок лицензии", lambda u, d, p: str(u.licence_year) if u.licence_year else "—"),
}


def _parse_date_bound(val: Optional[str], is_end: bool = False) -> Optional[datetime]:
    if not val:
        return None
    val = val.strip()
    try:
        if len(val) == 10:  # YYYY-MM-DD
            dt = datetime.strptime(val, "%Y-%m-%d")
            return dt.replace(hour=23, minute=59, second=59) if is_end else dt.replace(hour=0, minute=0, second=0)
        return datetime.fromisoformat(val)
    except Exception:
        return None


async def _fetch_report_data(
    db: AsyncSession,
    current_user: User,
    direction_id: Optional[int] = None,
    product_id: Optional[int] = None,
    manager_id: Optional[str] = None,
    period_from: Optional[str] = None,
    period_to: Optional[str] = None,
):
    query = select(University).where(University.is_active == True)
    if current_user.role == UserRole.manager:
        query = query.where(University.manager_id == current_user.id)
    elif manager_id:
        query = query.where(University.manager_id == manager_id)
        
    if direction_id:
        query = query.where(University.direction_id == direction_id)
    if product_id:
        query = query.where(University.product_id == product_id)

    dt_from = _parse_date_bound(period_from, is_end=False)
    dt_to = _parse_date_bound(period_to, is_end=True)

    if dt_from:
        history_subq = select(WorkflowHistory.university_id).where(WorkflowHistory.created_at >= dt_from)
        query = query.where(
            or_(
                University.updated_at >= dt_from,
                University.created_at >= dt_from,
                University.id.in_(history_subq),
            )
        )
    if dt_to:
        history_subq_to = select(WorkflowHistory.university_id).where(WorkflowHistory.created_at <= dt_to)
        query = query.where(
            or_(
                University.created_at <= dt_to,
                University.updated_at <= dt_to,
                University.id.in_(history_subq_to),
            )
        )

    unis = (await db.scalars(query)).all()
    
    # Кэшируем справочники
    dirs = {d.id: d for d in (await db.scalars(select(Direction))).all()}
    prods = {p.id: p for p in (await db.scalars(select(Product))).all()}
    stages = {s.order: s.name for s in (await db.scalars(select(WorkflowStage))).all()}

    return unis, dirs, prods, stages


@router.get("/generate", summary="Генерация отчёта в XLSX / PDF / JSON")
async def generate_report(
    format: str = Query("xlsx", pattern="^(xlsx|xls|pdf|json)$"),
    columns: Optional[str] = Query("name,direction,product,stage,manager"),
    direction_id: Optional[int] = None,
    product_id: Optional[int] = None,
    manager_id: Optional[str] = None,
    period_from: Optional[str] = None,
    period_to: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_manager),
):
    col_keys = [c.strip() for c in columns.split(",") if c.strip() in COLUMN_MAP]
    if not col_keys:
        col_keys = ["name", "direction", "product", "stage", "manager"]

    unis, dirs, prods, stages = await _fetch_report_data(
        db, current_user, direction_id, product_id, manager_id, period_from, period_to
    )

    rows = []
    for u in unis:
        d = dirs.get(u.direction_id)
        p = prods.get(u.product_id)
        row = {}
        for k in col_keys:
            if k == "stage":
                st_name = stages.get(u.current_stage_order, f"Этап {u.current_stage_order + 1}")
                row[k] = f"{u.current_stage_order + 1}. {st_name}"
            else:
                row[k] = COLUMN_MAP[k][1](u, d, p)
        rows.append(row)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")

    # 1. JSON
    period_str = f"с {period_from} по {period_to}" if (period_from or period_to) else "За всё время"
    if format == "json":
        data = {
            "title": "Отчёт по взаимодействию с вузами — ИТ Школа Ростелекома",
            "generated_at": datetime.now().isoformat(),
            "generated_by": current_user.full_name,
            "period": {
                "from": period_from,
                "to": period_to,
                "label": period_str,
            },
            "total": len(rows),
            "columns": [COLUMN_MAP[k][0] for k in col_keys],
            "items": rows,
        }
        return Response(
            content=json.dumps(data, ensure_ascii=False, indent=2),
            media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="report_{timestamp}.json"'}
        )

    # 2. XLSX / XLS
    if format in ("xlsx", "xls"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Взаимодействие с ВУЗами"

        # Стили
        header_fill = PatternFill(start_color="7700FF", end_color="7700FF", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        meta_font = Font(name="Calibri", size=9, italic=True, color="666666")
        regular_font = Font(name="Calibri", size=10)
        border_side = Side(style="thin", color="CCCCCC")
        thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        # Мета-информация о периоде в строке 1
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(col_keys))
        meta_cell = ws.cell(row=1, column=1)
        meta_cell.value = f"Отчёт по взаимодействию с вузами | Период: {period_str} | Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')} ({current_user.full_name})"
        meta_cell.font = meta_font

        # Заголовки таблицы в строке 2
        headers = [COLUMN_MAP[k][0] for k in col_keys]
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Строки данных начиная со строки 3
        for r_idx, r in enumerate(rows, start=3):
            for col_idx, k in enumerate(col_keys, start=1):
                cell = ws.cell(row=r_idx, column=col_idx, value=r[k])
                cell.font = regular_font
                cell.border = thin_border

        # Автоширина колонок
        for col in ws.columns:
            vals = [str(cell.value or "") for cell in col if cell.row > 1]
            max_len = max(len(v) for v in vals) if vals else 12
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        ext = "xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="report_{timestamp}.{ext}"'}
        )

    # 3. PDF
    if format == "pdf":
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=25, bottomMargin=25)
        story = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "Title",
            parent=styles["Heading1"],
            fontName=PDF_FONT,
            fontSize=16,
            textColor=colors.HexColor("#7700FF"),
            spaceAfter=8
        )
        meta_style = ParagraphStyle("Meta", fontName=PDF_FONT, fontSize=9, textColor=colors.HexColor("#555555"))
        body_style = ParagraphStyle("Body", fontName=PDF_FONT, fontSize=9)

        story.append(Paragraph("Отчёт по взаимодействию с вузами — ИТ Школа Ростелекома", title_style))
        story.append(Paragraph(f"<b>Период отчета:</b> {period_str} | <b>Дата выгрузки:</b> {datetime.now().strftime('%d.%m.%Y %H:%M')} | <b>Сформировал:</b> {current_user.full_name}", meta_style))
        story.append(Spacer(1, 12))

        table_data = [[COLUMN_MAP[k][0] for k in col_keys]]
        for r in rows:
            table_data.append([Paragraph(str(r[k]), body_style) for k in col_keys])

        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7700FF")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), PDF_FONT),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
        ]))
        story.append(t)
        doc.build(story)
        buf.seek(0)

        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="report_{timestamp}.pdf"'}
        )
